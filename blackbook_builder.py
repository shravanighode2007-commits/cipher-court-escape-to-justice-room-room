# blackbook_builder.py
# Master two-pass compiler and Word COM exporter for Cipher Court Project Blackbook

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from blackbook_helpers import FONT_NAME
from blackbook_content_mod1 import (
    add_front_matter, add_chapter_1, add_chapter_2, add_chapter_3, add_chapter_4, add_chapter_5
)
from blackbook_content_mod2 import (
    add_chapter_6, add_chapter_7, add_chapter_8, add_chapter_9, add_chapter_10, add_chapter_11_and_refs
)

def build_document(toc_data, tot_data, tof_data, target_docx_path):
    """Assembles the complete Word document structure."""
    doc = docx.Document()

    # ══════════════════════════════════════════════════════════
    # SECTION 1: FRONT MATTER (Roman Page Numbers)
    # ══════════════════════════════════════════════════════════
    sec1 = doc.sections[0]
    sec1.top_margin = Inches(1.0)
    sec1.bottom_margin = Inches(1.0)
    sec1.left_margin = Inches(1.0)
    sec1.right_margin = Inches(1.0)
    sec1.different_first_page_header_footer = True

    # Roman numerals for Front Matter
    sectPr1 = sec1._sectPr
    pgNumType1 = parse_xml(f'<w:pgNumType {nsdecls("w")} w:fmt="lowerRoman" w:start="1"/>')
    sectPr1.append(pgNumType1)

    # Centered page number in footer for pages 2+
    footer1 = sec1.footer
    f_p1 = footer1.paragraphs[0]
    f_p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    f_run1 = f_p1.add_run()
    fld1 = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="PAGE"/>')
    f_p1._p.append(fld1)

    # Build Front Matter
    add_front_matter(doc, toc_data, tot_data, tof_data)

    # ══════════════════════════════════════════════════════════
    # SECTION 2: CHAPTER 1 ONWARDS (Arabic Numbers Restart at 1)
    # ══════════════════════════════════════════════════════════
    sec2 = doc.add_section(WD_SECTION.NEW_PAGE)
    sec2.top_margin = Inches(1.0)
    sec2.bottom_margin = Inches(1.0)
    sec2.left_margin = Inches(1.0)
    sec2.right_margin = Inches(1.0)
    sec2.different_first_page_header_footer = False
    sec2.header.is_linked_to_previous = False
    sec2.footer.is_linked_to_previous = False

    # Arabic decimal starting at 1
    sectPr2 = sec2._sectPr
    pgNumType2 = parse_xml(f'<w:pgNumType {nsdecls("w")} w:fmt="decimal" w:start="1"/>')
    sectPr2.append(pgNumType2)

    # Centered Arabic page number in footer
    footer2 = sec2.footer
    f_p2 = footer2.paragraphs[0]
    f_p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld2 = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="PAGE"/>')
    f_p2._p.append(fld2)

    # Running Header
    header2 = sec2.header
    h_p2 = header2.paragraphs[0]
    h_p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h_run = h_p2.add_run("Modern Education Society’s The D. G. Ruparel College | Department of Computer Science")
    h_run.font.name = FONT_NAME
    h_run.font.size = Pt(8.5)
    h_run.font.italic = True
    h_run.font.color.rgb = RGBColor(120, 120, 120)

    # Assemble All Chapters
    print("  -> Adding Chapter 1 (Problem Identification & Feasibility Study)...")
    add_chapter_1(doc)

    print("  -> Adding Chapter 2 (Requirement Engineering)...")
    add_chapter_2(doc)

    print("  -> Adding Chapter 3 (SDLC Planning & Prototype Model)...")
    add_chapter_3(doc)

    print("  -> Adding Chapter 4 (System Modeling using UML - 1 Diagram Per Page)...")
    add_chapter_4(doc)

    print("  -> Adding Chapter 5 (System Architecture & Design)...")
    add_chapter_5(doc)

    print("  -> Adding Chapter 6 (Application Development)...")
    add_chapter_6(doc)

    print("  -> Adding Chapter 7 (Integration & System Testing)...")
    add_chapter_7(doc)

    print("  -> Adding Chapter 8 (Deployment & Hosting)...")
    add_chapter_8(doc)

    print("  -> Adding Chapter 9 (Performance & Security Testing)...")
    add_chapter_9(doc)

    print("  -> Adding Chapter 10 (Final Documentation / Results and Discussion)...")
    add_chapter_10(doc)

    print("  -> Adding Chapter 11 (Conclusion, References & Glossary)...")
    add_chapter_11_and_refs(doc)

    print(f"Saving compiled document to: {target_docx_path}")
    doc.save(target_docx_path)
    return doc

def extract_dynamic_page_numbers(docx_path):
    """Uses Word COM to determine the exact adjusted page number of each heading, table, and figure in Section 2."""
    import win32com.client as win32
    print("Connecting to Word COM to extract accurate dynamic page numbers...")
    word = win32.Dispatch("Word.Application")
    word.Visible = False
    wdoc = word.Documents.Open(os.path.abspath(docx_path))
    
    # Force layout calculation
    total_pages = wdoc.ComputeStatistics(2)
    print(f"Total pages rendered in Word: {total_pages}")

    page_map = {}
    sec2 = wdoc.Sections(2)
    for p in sec2.Range.Paragraphs:
        txt = p.Range.Text.strip()
        if txt and (
            txt.startswith("MODULE ") or 
            txt.startswith("CHAPTER ") or 
            txt.startswith("Table ") or 
            txt.startswith("Figure ") or
            txt.startswith("REFERENCES") or
            txt.startswith("GLOSSARY") or
            txt.startswith("Role and Responsibility") or
            any(txt.startswith(f"{i}.") for i in range(1, 12))
        ):
            pg = p.Range.Information(1) # wdActiveEndAdjustedPageNumber = 1 (restarts at 1 in Section 2!)
            if txt not in page_map:
                page_map[txt] = pg

    wdoc.Close(False)
    word.Quit()
    return page_map, total_pages

def build_full_blackbook():
    print("=================================================================")
    print("STARTING CIPHER COURT PROJECT BLACKBOOK COMPILATION")
    print("Candidate: Ms. Shravani Laxman Ghode (Roll No. 9022)")
    print("Modern Education Society’s The D. G. Ruparel College, Mumbai")
    print("=================================================================")

    # Initial estimated TOC, TOT, TOF mappings
    initial_toc = [
        ("-", "CERTIFICATE", "ii"),
        ("-", "ACKNOWLEDGEMENTS", "iii"),
        ("-", "DECLARATION", "iv"),
        ("-", "LIST OF ABBREVIATIONS", "v"),
        ("-", "TABLE OF CONTENTS", "vi"),
        ("-", "TABLE OF TABLES", "vii"),
        ("-", "TABLE OF FIGURES", "viii"),
        ("MODULE 1", "PROBLEM IDENTIFICATION, REQUIREMENT ENGINEERING & SYSTEM DESIGN PHASE", "1"),
        ("CHAPTER 1", "PROBLEM IDENTIFICATION & FEASIBILITY STUDY", "1"),
        ("-", "Role and Responsibility", "1"),
        ("1.1", "Background and Motivation", "1"),
        ("1.2", "Objectives", "2"),
        ("1.3", "Identification of a Real-World Problem (Industry / Social / Institutional)", "3"),
        ("1.4", "Problem Justification and Scope Definition", "3"),
        ("1.5", "Stakeholder Identification & Beneficiary Analysis", "3"),
        ("1.6", "Feasibility Analysis (Technical, Economic, Operational, Legal, Schedule)", "4"),
        ("CHAPTER 2", "REQUIREMENT ENGINEERING", "5"),
        ("2.1", "Functional Requirements Specification", "5"),
        ("2.2", "Non-Functional Requirements Specification", "7"),
        ("2.3", "Use-Case Analysis & Actor Profiles", "8"),
        ("2.4", "Requirement Prioritization (MoSCoW Matrix)", "9"),
        ("2.5", "Constraints and Assumptions", "9"),
        ("2.6", "Preliminary Product Description & Technology Stack Survey", "10"),
        ("2.7", "Conceptual Models (Literature Review of Existing Serious Game Systems)", "11"),
        ("CHAPTER 3", "SOFTWARE DEVELOPMENT LIFE CYCLE (SDLC) PLANNING", "13"),
        ("3.1", "Selection of SDLC Model (Prototype Model)", "13"),
        ("3.2", "Work Breakdown Structure (WBS)", "13"),
        ("3.3", "Project Timeline & Scheduling (Gantt Chart & Sprint Roadmap)", "14"),
        ("3.4", "Resource Planning (Hardware & Software Specifications)", "16"),
        ("CHAPTER 4", "SYSTEM MODELING USING UML", "17"),
        ("4.1", "Event Table (System Logic & Event Table)", "17"),
        ("4.2", "Class Diagram", "18"),
        ("4.3", "Object Diagram", "19"),
        ("4.4", "Use Case Diagram", "20"),
        ("4.5", "Sequence Diagram", "21"),
        ("4.6", "Activity Diagram", "22"),
        ("4.7", "Component Diagram", "23"),
        ("4.8", "Deployment Diagram", "24"),
        ("4.9", "Faculty Feedback: Frontend Wireframe Prototype & Navigation Flow", "25"),
        ("4.10", "Faculty Feedback: Backend Architecture & Implementation Flow", "26"),
        ("CHAPTER 5", "SYSTEM ARCHITECTURE DESIGN", "27"),
        ("5.1", "Frontend Prototype Architecture & Component Hierarchy", "27"),
        ("5.2", "Backend Architecture & Game Logic", "28"),
        ("5.3", "Database Schema Design & Data Dictionaries", "28"),
        ("5.4", "API Structure, Endpoints & JSON Contracts", "32"),
        ("5.5", "Security Considerations & Granular Access Control Policies", "33"),
        ("5.6", "Procedural Design (Generic Puzzle Logic Flowchart & Algorithms)", "34"),
        ("MODULE 2", "IMPLEMENTATION, TESTING, DEPLOYMENT & EVALUATION PHASE", "36"),
        ("CHAPTER 6", "APPLICATION DEVELOPMENT & PRACTICAL IMPLEMENTATION", "36"),
        ("6.1", "Frontend Implementation Architecture & Component Flow", "36"),
        ("6.2", "Backend Implementation Architecture & Business Logic", "45"),
        ("6.3", "Integration & Environment Configuration", "52"),
        ("CHAPTER 7", "INTEGRATION & SYSTEM TESTING", "54"),
        ("7.1", "Testing Approach & Quality Assurance Framework", "54"),
        ("7.2", "Unit Testing", "54"),
        ("7.3", "Black-Box Testing", "54"),
        ("7.4", "Integration Testing", "54"),
        ("7.5", "Beta Testing & Usability Evaluation", "54"),
        ("7.6", "Comprehensive Test Case Matrix (Unit, Integration & System Test Tables)", "55"),
        ("7.7", "Bug Tracking & Defect Management Log", "59"),
        ("CHAPTER 8", "DEPLOYMENT & HOSTING", "60"),
        ("8.1", "Multi-Platform Hosting Architecture & Deployment Topology", "60"),
        ("8.2", "Build Pipeline (Unity Standalone Windows x86_64 Build)", "61"),
        ("8.3", "Live Online WebGL Deployment & Cloudflare Global Edge Tunneling", "61"),
        ("8.4", "Dual Backend Server Configuration & Secure Environment Variables (.env)", "61"),
        ("8.5", "Version Control using GitHub & Project Directory Tree", "62"),
        ("8.6", "Local & Cloud Database Server Setup & Containerization", "62"),
        ("8.7", "1-Click Operational Launchers & Deployment Verification Checklist", "63"),
        ("CHAPTER 9", "PERFORMANCE & SECURITY TESTING", "64"),
        ("9.1", "Basic Load Testing & Latency Benchmarks (60 FPS & <50ms Target)", "64"),
        ("9.2", "Input Validation Checks & Sanitization", "64"),
        ("9.3", "Security Validation & Cryptographic Protection", "65"),
        ("CHAPTER 10", "TESTING, EVALUATION, RESULTS & FUTURE ROADMAP", "66"),
        ("10.1", "Comprehensive Test Execution & Evidence", "66"),
        ("10.2", "Performance, Load & Security Analysis", "73"),
        ("10.3", "Step-by-Step User Scenarios & Operational Walkthrough", "77"),
        ("10.4", "Application Visual Verification & Interface Analysis", "81"),
        ("10.5", "Source Code Documentation & Algorithmic Explanations", "85"),
        ("10.6", "Limitations & Edge Cases Encountered", "89"),
        ("10.7", "Strategic Roadmap & Future Enhancements", "90"),
        ("10.8", "Final Project Summary & Academic Contribution", "91"),
        ("CHAPTER 11", "CONCLUSION", "92"),
        ("11.1", "Significance of the System (Concluding Remarks)", "92"),
        ("11.2", "Limitations of the System", "92"),
        ("11.3", "Future Scope of the Project", "93"),
        ("-", "REFERENCES", "94"),
        ("-", "GLOSSARY", "96")
    ]

    initial_tot = [
        ("Table 1.1", "Role and Responsibility Across Project Phases", "1"),
        ("Table 2.1", "Functional Requirements Specification Table", "5"),
        ("Table 2.2", "Non-Functional Requirements Specification Table", "7"),
        ("Table 2.3", "Requirement Prioritization (MoSCoW Matrix)", "9"),
        ("Table 2.4", "Comparative Analysis Matrix of Existing Cybersecurity Learning Systems", "11"),
        ("Table 3.1", "Work Breakdown Structure (WBS) & Prototype Deliverables", "13"),
        ("Table 4.1", "System Logic & Event Table for Cipher Court Gameplay", "17"),
        ("Table 5.1", "Level-Wise Frontend Prototype Architecture & Component Specification Matrix", "28"),
        ("Table 5.2", "PLAYER (Player Identity & Profiles) Data Dictionary", "29"),
        ("Table 5.3", "LEVEL (Stage Configurations) Data Dictionary", "30"),
        ("Table 5.4", "GAME_PROGRESS (Player State) Data Dictionary", "30"),
        ("Table 5.5", "CLUE (Forensic Clues) Data Dictionary", "31"),
        ("Table 5.6", "PUZZLE_ATTEMPT (Submission Audit Logs) Data Dictionary", "31"),
        ("Table 5.7", "EVIDENCE_RECORD (Collected Evidence Registry) Data Dictionary", "32"),
        ("Table 5.8", "REST API Endpoints & JSON Contracts", "33"),
        ("Table 6.1", "Frontend UI Component Architecture & Design Token Registry", "38"),
        ("Table 6.2", "Frontend Client-Side Event Handlers & Input Mappings", "39"),
        ("Table 6.3", "Client-Side Input Sanitization & Defensive Validation Protocols", "40"),
        ("Table 6.4", "Backend REST API Endpoints & Request/Response Specification", "47"),
        ("Table 6.5", "HTTP Status Codes & Error Handling Matrix", "48"),
        ("Table 6.6", "Environment Variables & Security Configuration Matrix", "53"),
        ("Table 6.7", "Relational Database Indexing Strategy & Query Optimization Matrix", "54"),
        ("Table 7.1", "Unit Testing Test Cases Log", "55"),
        ("Table 7.2", "Integration Testing Test Cases Log", "56"),
        ("Table 7.3", "System & Black-Box Testing Test Cases Log", "57"),
        ("Table 7.4", "Bug Tracking & Defect Management Log", "59"),
        ("Table 8.1", "Multi-Platform Deployment & Hosting Architecture Matrix", "60"),
        ("Table 8.2", "Server Configuration, Ports & Environment Variables (.env)", "61"),
        ("Table 8.3", "End-to-End Deployment Verification & Operational Health Checklist", "63"),
        ("Table 10.1", "Comprehensive Unit Testing Matrix (Algorithmic Logic & Component Bounds)", "66"),
        ("Table 10.2", "Integration Testing & REST API Verification Matrix", "67"),
        ("Table 10.3", "Regression Testing Matrix (Bug Fix Verifications Across Releases)", "68"),
        ("Table 10.4", "End-to-End (E2E) User Journey Testing Matrix", "69"),
        ("Table 10.5", "Empirical Performance & Latency Benchmarks Across Hardware Tiers", "73"),
        ("Table 10.6", "Backend Concurrency & Stress Ingestion Benchmarks", "74"),
        ("Table 10.7", "Security Audit & Vulnerability Assessment Findings", "75")
    ]

    initial_tof = [
        ("Figure 3.1", "Project Schedule and Prototype Development Gantt Chart Sem-V 2026-27", "14"),
        ("Figure 4.1", "System Logic & Event Table for Cipher Court Gameplay", "17"),
        ("Figure 4.2", "Class Diagram Mapping System Entities and Object Relationships", "18"),
        ("Figure 4.3", "Object Diagram Illustrating Runtime Instances during Active Gameplay", "19"),
        ("Figure 4.4", "Use Case Diagram for Cipher Court System Actors and Boundaries", "20"),
        ("Figure 4.5", "Sequence Diagram for Gameplay, Interaction, and Puzzle Validation Lifecycle", "21"),
        ("Figure 4.6", "Activity Diagram for Level Progression and Puzzle Verification Workflow", "22"),
        ("Figure 4.7", "Component Diagram Showing Subsystem Interconnections and Service Decoupling", "23"),
        ("Figure 4.8", "Deployment Diagram Across Client Workstation and Backend Database Server", "24"),
        ("Figure 4.9", "Frontend Navigation Flow and Screen State Hierarchy", "25"),
        ("Figure 4.10", "Backend Data Flow and System Processing Pipeline", "26"),
        ("Figure 5.1", "Level 0 Main Menu Prototype Interface — Cinematic Courtroom Gateway (MainMenu.unity)", "27"),
        ("Figure 5.2", "Level 1 Crime Scene Prototype — Reticle Clue Interaction Prompt Active (Level1_Investigation.unity)", "28"),
        ("Figure 5.3", "Level 2 Cyber Forensics Terminal & Password Decryption Prototype (Level2_Computer.unity)", "29"),
        ("Figure 5.4", "Level 3 Steganographic Caesar Cipher Decoding Prototype (Level3_Evidence.unity)", "30"),
        ("Figure 5.5", "Level 4 Underground Archival Vault Records Prototype (Level4_Archive.unity)", "31"),
        ("Figure 5.6", "Level 5 Case Files & Sworn Affidavits Deliberation Prototype (Level5_Statements.unity)", "32"),
        ("Figure 5.7", "Level 6 Grand Courtroom Cross-Examination & Perjury Detection Prototype (Level6_Courtroom.unity)", "33"),
        ("Figure 5.8", "Level 7 Victory Debriefing Prototype — Case Solved, Justice Served (GameComplete.unity)", "34"),
        ("Figure 5.9", "Proposed Entity-Relationship (ER) Diagram for Cipher Court Database Schema", "35"),
        ("Figure 5.10", "Generic Puzzle Business Logic and Attempt Validation Flowchart", "36"),
        ("Figure 6.1", "Frontend Authentication & Investigator Profile Gateway Interface (MainMenu.unity)", "41"),
        ("Figure 6.2", "Frontend Navigation Hierarchy and Live Telemetry Dashboard", "42"),
        ("Figure 6.3", "Frontend Level 1 Forensic Crime Scene with Reticle Clue Interaction Prompt Active", "43"),
        ("Figure 6.4", "Frontend Level 2 Cyber Forensics Terminal and Password Authentication Modal", "44"),
        ("Figure 6.5", "Frontend Level 3 Steganographic Caesar Cipher Decoding Interface (Ciphertext: MXVWLFH)", "45"),
        ("Figure 6.6", "Frontend Archival Evidence Vault with Interactive Numeric Keypad Safe Interface", "46"),
        ("Figure 6.7", "Frontend Level 6 Grand Courtroom Witness Cross-Examination Interface (Level6_Courtroom.unity)", "47"),
        ("Figure 6.8", "Frontend Game Complete Victory Screen — Case Solved, Justice Served (GameComplete.unity)", "48"),
        ("Figure 6.9", "Backend Interactive OpenAPI / Swagger UI Specification (/docs)", "50"),
        ("Figure 6.10", "Backend Postman Automated REST Test Collection Runner & Validation Report", "51"),
        ("Figure 6.11", "Backend Database Management Console & Forensic Audit Trail Query (cipher_court.db)", "52"),
        ("Figure 6.12", "Backend Threaded Daemon Server Log Terminal & Event Ingestion Pipeline (server.py)", "53"),
        ("Figure 8.1", "Complete Architecture Overview and Multi-Platform Deployment Infrastructure", "60")
    ]

    # PASS 1: Build draft document
    draft_path = os.path.abspath("e:/vyt/Cipher_Court_Project_Report_Draft.docx")
    print("\n--- PASS 1: Generating Document Draft ---")
    build_document(initial_toc, initial_tot, initial_tof, draft_path)

    # Query Word COM for exact page positions
    page_map, total_pages = extract_dynamic_page_numbers(draft_path)

    def find_page(prefix, default_val):
        clean_prefix = prefix.strip()
        # Look for exact match or prefix followed by colon or space
        for k, v in page_map.items():
            k_strip = k.strip()
            if k_strip == clean_prefix:
                return str(v)
            if k_strip.startswith(clean_prefix + ":") or k_strip.startswith(clean_prefix + " "):
                return str(v)
        for k, v in page_map.items():
            if k.strip().startswith(clean_prefix):
                return str(v)
        return str(default_val)

    # Update TOC, TOT, TOF with accurate dynamic page numbers
    updated_toc = []
    for num, desc, default_p in initial_toc:
        if num == "-":
            search_key = desc
        elif "CHAPTER" in num:
            search_key = f"{num}"
        else:
            search_key = f"{num} "
        exact_p = find_page(search_key, default_p)
        updated_toc.append((num, desc, exact_p))

    updated_tot = []
    for t_no, t_desc, default_p in initial_tot:
        exact_p = find_page(t_no, default_p)
        updated_tot.append((t_no, t_desc, exact_p))

    updated_tof = []
    for f_no, f_desc, default_p in initial_tof:
        exact_p = find_page(f_no, default_p)
        updated_tof.append((f_no, f_desc, exact_p))

    # PASS 2: Build final document with verified exact page numbers
    final_docx = os.path.abspath("e:/vyt/Cipher_Court_Project_Report.docx")
    print("\n--- PASS 2: Generating Final Document with Dynamic Page Numbers ---")
    build_document(updated_toc, updated_tot, updated_tof, final_docx)

    # Export to PDF via Word COM
    try:
        import win32com.client as win32
        print("\n--- EXPORTING TO PDF VIA WORD COM ---")
        word = win32.Dispatch("Word.Application")
        word.Visible = False
        wdoc = word.Documents.Open(final_docx)
        final_pages = wdoc.ComputeStatistics(2)
        pdf_path = os.path.abspath("e:/vyt/Cipher_Court_Project_Report.pdf")
        wdoc.SaveAs(pdf_path, FileFormat=17) # 17 = wdFormatPDF
        wdoc.Close(False)
        word.Quit()
        print(f"=================================================================")
        print(f"SUCCESSFULLY GENERATED CIPHER COURT PROJECT BLACKBOOK!")
        print(f"Total Page Count: {final_pages} Pages")
        print(f"Word Document: {final_docx}")
        print(f"PDF Document: {pdf_path}")
        print(f"=================================================================")
        return final_pages
    except Exception as e:
        print(f"PDF Export Notice: {e}")
        return total_pages

if __name__ == "__main__":
    build_full_blackbook()
