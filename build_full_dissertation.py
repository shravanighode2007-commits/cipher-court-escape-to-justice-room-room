import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from report_data_frontmatter import add_front_matter
from report_data_ch1_ch2 import add_chapter_1_and_2
from report_data_ch3_ch4 import add_chapter_3_and_4
from report_data_ch5_code import add_chapter_5
from report_data_ch6_tests import add_chapter_6
from report_data_ch7_ch8_refs import add_chapter_7_8_and_refs
from report_blackbook_data import add_database_chapter, add_proposal_and_srs

def generate_dissertation():
    print("Initializing Word Document structure...")
    doc = docx.Document()

    # 1.0 inch margins all around
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)
        s.different_first_page_header_footer = True
        
        # Header (Pages 2+)
        header = s.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        hrun = hp.add_run("CIPHER COURT: Escape to Justice – B.Tech Major Project Dissertation Report")
        hrun.font.name = "Times New Roman"
        hrun.font.size = Pt(8.5)
        hrun.font.italic = True
        hrun.font.color.rgb = RGBColor(90, 90, 90)

    # Color Tokens
    NAVY = RGBColor(16, 44, 87)       # #102C57
    STEEL = RGBColor(53, 89, 142)     # #35598E
    CHARCOAL = RGBColor(35, 35, 35)   # #232323
    GRAY = RGBColor(100, 100, 100)
    DARK_CODE = RGBColor(20, 30, 45)

    def set_cell_background(cell, fill_hex):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        tcPr.append(shd)

    def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
        tcPr.append(tcMar)

    def set_table_borders(table, color="D3D3D3"):
        tblPr = table._tbl.tblPr
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
            f'<w:left w:val="none"/>'
            f'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="{color}"/>'
            f'<w:right w:val="none"/>'
            f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr.append(borders)

    def add_chapter_heading(title, number_str):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(26)
        p.paragraph_format.space_after = Pt(14)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(f"CHAPTER {number_str}: {title.upper()}")
        run.font.name = "Times New Roman"
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = NAVY
        return p

    def add_sec_heading(title, sec_str):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(f"{sec_str} {title}")
        run.font.name = "Times New Roman"
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = STEEL
        return p

    def add_subsec_heading(title, subsec_str):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(f"{subsec_str} {title}")
        run.font.name = "Times New Roman"
        run.font.size = Pt(11.5)
        run.font.bold = True
        run.font.color.rgb = CHARCOAL
        return p

    def add_body_p(text, bold_prefix=None, space_after=9.5):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2.5)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.4
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if bold_prefix:
            br = p.add_run(bold_prefix + " ")
            br.font.name = "Times New Roman"
            br.font.size = Pt(11)
            br.font.bold = True
            br.font.color.rgb = CHARCOAL
        run = p.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(11)
        run.font.color.rgb = CHARCOAL
        return p

    def add_bullet_p(text, bold_prefix=None):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.3
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if bold_prefix:
            br = p.add_run(bold_prefix + " ")
            br.font.name = "Times New Roman"
            br.font.size = Pt(10.5)
            br.font.bold = True
            br.font.color.rgb = CHARCOAL
        run = p.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(10.5)
        run.font.color.rgb = CHARCOAL
        return p

    def add_code_block(code_text, caption=None):
        if caption:
            cp = doc.add_paragraph()
            cp.paragraph_format.space_before = Pt(12)
            cp.paragraph_format.space_after = Pt(4)
            cp.paragraph_format.keep_with_next = True
            crun = cp.add_run(f"Listing: {caption}")
            crun.font.name = "Times New Roman"
            crun.font.size = Pt(9.5)
            crun.font.bold = True
            crun.font.color.rgb = STEEL

        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.5)
        set_cell_background(cell, "F4F6F9")
        set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
        
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="102C57"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
        tcPr.append(borders)

        cp = cell.paragraphs[0]
        cp.paragraph_format.space_before = Pt(3)
        cp.paragraph_format.space_after = Pt(3)
        cp.paragraph_format.line_spacing = 1.12
        run = cp.add_run(code_text.strip())
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
        run.font.color.rgb = DARK_CODE
        
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    def add_image_figure(img_path, caption, width=Inches(5.6)):
        if os.path.exists(img_path):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run()
            run.add_picture(img_path, width=width)
            
            cp = doc.add_paragraph()
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cp.paragraph_format.space_before = Pt(2)
            cp.paragraph_format.space_after = Pt(10)
            crun = cp.add_run(caption)
            crun.font.name = "Times New Roman"
            crun.font.size = Pt(9.5)
            crun.font.italic = True
            crun.font.color.rgb = GRAY
        else:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(f"[{caption} - Image: {os.path.basename(img_path)}]")
            run.font.italic = True

    helpers = {
        'NAVY': NAVY,
        'STEEL': STEEL,
        'CHARCOAL': CHARCOAL,
        'GRAY': GRAY,
        'DARK_CODE': DARK_CODE,
        'set_cell_background': set_cell_background,
        'set_cell_margins': set_cell_margins,
        'set_table_borders': set_table_borders,
        'add_chapter_heading': add_chapter_heading,
        'add_sec_heading': add_sec_heading,
        'add_subsec_heading': add_subsec_heading,
        'add_body_p': add_body_p,
        'add_bullet_p': add_bullet_p,
        'add_code_block': add_code_block,
        'add_image_figure': add_image_figure,
    }

    print("Building Front Matter...")
    add_front_matter(doc, helpers)

    print("Building Chapter 1 & 2 (Proposal, Motivation & Literature Survey)...")
    add_chapter_1_and_2(doc, helpers)
    add_proposal_and_srs(doc, helpers)

    print("Building Chapter 3 & 4 (SRS, System Architecture & UML Models)...")
    add_chapter_3_and_4(doc, helpers)

    print("Building Chapter 5 (Database & Data Persistence Design)...")
    add_database_chapter(doc, helpers)

    print("Building Chapter 6 (Core Logic Code & Implementation)...")
    add_chapter_5(doc, helpers)

    print("Building Chapter 7 (Software Testing & QA with 45 Test Cases)...")
    add_chapter_6(doc, helpers)

    print("Building Chapter 8, 9, 10, References & Appendices (Frontend Screenshots & Evaluation)...")
    add_chapter_7_8_and_refs(doc, helpers)

    out_docx = os.path.abspath("Cipher_Court_Project_Report.docx")
    print(f"Saving compiled document to: {out_docx}")
    doc.save(out_docx)
    print("Word document generated successfully!")

    # Check page count and export to PDF via Word COM
    try:
        import win32com.client as win32
        print("Opening Word COM to compute statistics and export PDF...")
        word = win32.Dispatch("Word.Application")
        word.Visible = False
        wdoc = word.Documents.Open(out_docx)
        page_count = wdoc.ComputeStatistics(2) # 2 = wdStatisticPages
        print(f"==================================================")
        print(f"TOTAL PAGE COUNT IN MS WORD: {page_count} PAGES")
        print(f"==================================================")
        
        pdf_path = os.path.abspath("Cipher_Court_Project_Report.pdf")
        wdoc.SaveAs(pdf_path, FileFormat=17) # 17 = wdFormatPDF
        print(f"PDF exported successfully to: {pdf_path}")
        wdoc.Close(False)
        word.Quit()
        return page_count
    except Exception as e:
        print(f"Note: Word COM conversion notice: {e}")
        return -1

if __name__ == "__main__":
    generate_dissertation()
