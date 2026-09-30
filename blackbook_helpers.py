# blackbook_helpers.py
# Universal styling, XML manipulation, and formatting helper functions for Cipher Court Project Blackbook

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

FONT_NAME = "Times New Roman"

def set_cell_background(cell, fill_hex):
    """Sets background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    """Sets internal padding for a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="B0B0B0"):
    """Applies neat, academic-style borders to a table."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def set_table_grid_borders(table, color="C0C0C0"):
    """Applies grid borders to all cells for formal matrices and logs."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_module_banner(doc, module_num, module_title):
    """Adds Module Title Banner matching university syllabus copy."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(28)
    p.paragraph_format.space_after = Pt(14)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"MODULE {module_num}\n{module_title.upper()}")
    run.font.name = FONT_NAME
    run.font.size = Pt(16)
    run.font.bold = True
    return p

def add_chapter_title(doc, chapter_num, chapter_title):
    """Adds Chapter Title in Times New Roman, 16 pt, Bold."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(14)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"CHAPTER {chapter_num}\n\n{chapter_title.upper()}")
    run.font.name = FONT_NAME
    run.font.size = Pt(16)
    run.font.bold = True
    return p

def add_heading(doc, text, space_before=16, space_after=6):
    """Adds Section Heading in Times New Roman, 14 pt, Bold."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = FONT_NAME
    run.font.size = Pt(14)
    run.font.bold = True
    return p

def add_subheading(doc, text, space_before=12, space_after=4):
    """Adds Subheading in Times New Roman, 14 pt, Bold."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = FONT_NAME
    run.font.size = Pt(14)
    run.font.bold = True
    return p

def add_body_paragraph(doc, text, bold_prefix=None, space_after=6, line_spacing=1.15):
    """Adds Body Text in Times New Roman, 12 pt, Regular, Justified."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_prefix:
        r_pre = p.add_run(bold_prefix + " ")
        r_pre.font.name = FONT_NAME
        r_pre.font.size = Pt(12)
        r_pre.font.bold = True
    r = p.add_run(text)
    r.font.name = FONT_NAME
    r.font.size = Pt(12)
    return p

def add_bullet_point(doc, text, bold_prefix=None):
    """Adds Bullet Point in Times New Roman, 12 pt."""
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_prefix:
        r_pre = p.add_run(bold_prefix + " ")
        r_pre.font.name = FONT_NAME
        r_pre.font.size = Pt(12)
        r_pre.font.bold = True
    r = p.add_run(text)
    r.font.name = FONT_NAME
    r.font.size = Pt(12)
    return p

def add_image_diagram(doc, img_path, caption_text, width=Inches(5.8), space_before=10, space_after=8):
    """Embeds an image figure with centered italic caption in Times New Roman."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if os.path.exists(img_path):
        r = p.add_run()
        r.add_picture(img_path, width=width)
    else:
        r = p.add_run(f"[{caption_text} - Image Missing: {os.path.basename(img_path)}]")
        r.font.name = FONT_NAME
        r.font.size = Pt(11)
        r.font.italic = True

    cp = doc.add_paragraph()
    cp.paragraph_format.space_before = Pt(2)
    cp.paragraph_format.space_after = Pt(space_after)
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c_run = cp.add_run(caption_text)
    c_run.font.name = FONT_NAME
    c_run.font.size = Pt(11)
    c_run.font.italic = True
    return cp

def add_screenshot_placeholder(doc, placeholder_tag, caption_text, img_path=None, width=Inches(5.6), space_before=8, space_after=8):
    """
    Renders an academic screenshot placeholder matching thesis formatting guidelines.
    Includes the explicit placeholder callout: [Insert Screenshot: ...]
    Embeds image if present; otherwise renders a styled academic placeholder box.
    Adds the formal italic figure caption.
    """
    if img_path and os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.paragraph_format.space_before = Pt(space_before)
        p_img.paragraph_format.space_after = Pt(4)
        p_img.paragraph_format.keep_with_next = True
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_img = p_img.add_run()
        r_img.add_picture(img_path, width=width)
    else:
        p_tag = doc.add_paragraph()
        p_tag.paragraph_format.space_before = Pt(space_before)
        p_tag.paragraph_format.space_after = Pt(3)
        p_tag.paragraph_format.keep_with_next = True
        p_tag.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_tag = p_tag.add_run(placeholder_tag)
        r_tag.font.name = FONT_NAME
        r_tag.font.size = Pt(10.5)
        r_tag.font.bold = True
        r_tag.font.italic = True
        r_tag.font.color.rgb = RGBColor(16, 44, 87)

        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        cell = tbl.cell(0, 0)
        cell.width = width
        set_cell_background(cell, "F4F6F9")
        set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
        
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="12" w:space="0" w:color="35598E"/><w:top w:val="single" w:sz="6" w:space="0" w:color="CCCCCC"/><w:right w:val="single" w:sz="6" w:space="0" w:color="CCCCCC"/><w:bottom w:val="single" w:sz="12" w:space="0" w:color="35598E"/></w:tcBorders>')
        tcPr.append(borders)

        cp = cell.paragraphs[0]
        cp.paragraph_format.space_before = Pt(8)
        cp.paragraph_format.space_after = Pt(8)
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cp.add_run(f"SCREENSHOT DEMONSTRATION ARTIFACT\n{placeholder_tag}\n[High-Resolution Terminal / GUI Capture Rendered in Active Test Environment]")
        run.font.name = FONT_NAME
        run.font.size = Pt(10)
        run.font.italic = True
        run.font.color.rgb = RGBColor(70, 80, 95)

    cp = doc.add_paragraph()
    cp.paragraph_format.space_before = Pt(2)
    cp.paragraph_format.space_after = Pt(space_after)
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c_run = cp.add_run(caption_text)
    c_run.font.name = FONT_NAME
    c_run.font.size = Pt(11)
    c_run.font.italic = True
    return cp


def add_single_page_diagram(doc, img_path, fig_id, fig_title, theory_paragraphs, width=Inches(5.8)):
    """
    Enforces the Layout Rule: Strictly one diagram per page!
    Includes Heading, Theoretical Explanation, Image, and Caption, followed by Page Break.
    """
    add_heading(doc, f"{fig_id}: {fig_title}")
    for p_text in theory_paragraphs:
        add_body_paragraph(doc, p_text)
    add_image_diagram(doc, img_path, f"Figure {fig_id.split()[0]}: {fig_title}", width=width)
    doc.add_page_break()

def add_code_block(doc, code_str, caption=None):
    """Renders formatted source code listings in single cell bordered box."""
    if caption:
        cp = doc.add_paragraph()
        cp.paragraph_format.space_before = Pt(10)
        cp.paragraph_format.space_after = Pt(3)
        cp.paragraph_format.keep_with_next = True
        crun = cp.add_run(caption)
        crun.font.name = FONT_NAME
        crun.font.size = Pt(11)
        crun.font.bold = True

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F7F9FA")
    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="18" w:space="0" w:color="003366"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(borders)

    cp = cell.paragraphs[0]
    cp.paragraph_format.space_before = Pt(2)
    cp.paragraph_format.space_after = Pt(2)
    cp.paragraph_format.line_spacing = 1.05
    run = cp.add_run(code_str.strip())
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(20, 30, 45)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)
