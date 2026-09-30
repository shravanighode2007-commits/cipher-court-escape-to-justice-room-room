import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

doc = docx.Document()
sec1 = doc.sections[0]
sec1.top_margin = Inches(1.0)
sec1.bottom_margin = Inches(1.0)

# Add front matter page
p = doc.add_paragraph('Front Matter Page i')
doc.add_page_break()
doc.add_paragraph('Front Matter Page ii')

# Add section break
sec2 = doc.add_section(WD_SECTION.NEW_PAGE)
sec2.header.is_linked_to_previous = False
sec2.footer.is_linked_to_previous = False

# Restart page numbering at 1
sectPr2 = sec2._sectPr
pgNumType = parse_xml(f'<w:pgNumType {nsdecls("w")} w:fmt="decimal" w:start="1"/>')
sectPr2.append(pgNumType)

# Add page number field to footer
footer_p = sec2.footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
fldSimple = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="PAGE"/>')
footer_p._p.append(fldSimple)

# Add Chapter 1
p1 = doc.add_paragraph('CHAPTER 1: Introduction')
doc.add_page_break()
p2 = doc.add_paragraph('Chapter 1 content on page 2')

test_path = 'e:/vyt/test_num.docx'
doc.save(test_path)
print('Saved test docx')

import win32com.client as win32
word = win32.Dispatch('Word.Application')
wdoc = word.Documents.Open(test_path)
pages = wdoc.ComputeStatistics(2)
print('Total pages:', pages)
wdoc.Close(False)
word.Quit()
print('Success!')
