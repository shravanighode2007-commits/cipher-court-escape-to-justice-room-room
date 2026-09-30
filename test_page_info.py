import docx
import win32com.client as win32

doc = docx.Document()
p1 = doc.add_paragraph('Heading 1')
doc.add_page_break()
p2 = doc.add_paragraph('Heading 2')
doc.save('e:/vyt/test_page.docx')

word = win32.Dispatch('Word.Application')
wdoc = word.Documents.Open('e:/vyt/test_page.docx')
for i, p in enumerate(wdoc.Paragraphs):
    txt = p.Range.Text.strip()
    if txt:
        page = p.Range.Information(3) # wdActiveEndPageNumber = 3
        print(f"P: '{txt}' -> Page {page}")
wdoc.Close(False)
word.Quit()
